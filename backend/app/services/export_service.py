import os
from bson import ObjectId
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from app.core.database import get_db
from app.core.config import settings


def _normalize_language(value: str) -> str:
    raw = (value or "").strip().lower()

    if raw in {"assamese", "as"}:
        return "as"

    if raw in {"hindi", "hi"}:
        return "hi"

    if raw in {"mix", "mixed", "multilingual", "mixed-language"}:
        return "mix"

    return "en"


def _get_labels(language: str):
    labels = {
        "en": {
            "transcript": "Transcript",
            "summary": "Summary",
            "discussion": "Discussion Points",
            "decisions": "Decisions",
            "actions": "Action Items",
            "next_steps": "Next Steps",
            "deadline": "Deadline",
        },
        "hi": {
            "transcript": "प्रतिलेख",
            "summary": "सारांश",
            "discussion": "चर्चा बिंदु",
            "decisions": "निर्णय",
            "actions": "कार्य बिंदु",
            "next_steps": "अगले कदम",
            "deadline": "समयसीमा",
        },
        "as": {
            "transcript": "প্ৰতিলিপি",
            "summary": "সাৰাংশ",
            "discussion": "আলোচনাৰ বিষয়সমূহ",
            "decisions": "সিদ্ধান্তসমূহ",
            "actions": "কৰণীয় কামসমূহ",
            "next_steps": "পৰৱৰ্তী পদক্ষেপসমূহ",
            "deadline": "সময়সীমা",
        },
        "mix": {
            "transcript": "Transcript / প্ৰতিলিপি / प्रतिलेख",
            "summary": "Summary / সাৰাংশ / सारांश",
            "discussion": "Discussion Points / আলোচনাৰ বিষয় / चर्चा बिंदु",
            "decisions": "Decisions / সিদ্ধান্ত / निर्णय",
            "actions": "Action Items / কৰণীয় / कार्य बिंदु",
            "next_steps": "Next Steps / পৰৱৰ্তী পদক্ষেপ / अगले कदम",
            "deadline": "Deadline / সময়সীমা / समयसीमा",
        },
    }

    return labels.get(language, labels["en"])


def _apply_run_font(run, font_name: str = "Nirmala UI", font_size: int | None = None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    run._element.rPr.rFonts.set(qn("w:cs"), font_name)

    if font_size is not None:
        run.font.size = Pt(font_size)


def _set_paragraph_font(paragraph, font_name: str = "Nirmala UI", font_size: int | None = None):
    if not paragraph.runs:
        run = paragraph.add_run("")
        _apply_run_font(run, font_name, font_size)
        return

    for run in paragraph.runs:
        _apply_run_font(run, font_name, font_size)


def _set_document_default_font(doc: Document, font_name: str = "Nirmala UI", font_size: int = 11):
    normal_style = doc.styles["Normal"]
    normal_style.font.name = font_name
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    normal_style._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    normal_style._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    normal_style._element.rPr.rFonts.set(qn("w:cs"), font_name)
    normal_style.font.size = Pt(font_size)

    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet"]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = font_name
            style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
            style._element.rPr.rFonts.set(qn("w:ascii"), font_name)
            style._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
            style._element.rPr.rFonts.set(qn("w:cs"), font_name)


async def export_minutes_docx(meeting_id: str):
    db = get_db()

    minutes = await db.minutes_drafts.find_one({"meeting_id": meeting_id})
    meeting = await db.meetings.find_one({"_id": ObjectId(meeting_id)})

    language = _normalize_language(
        (minutes or {}).get("output_language")
        or (meeting or {}).get("language")
    )

    labels = _get_labels(language)

    os.makedirs(settings.STORAGE_EXPORTS, exist_ok=True)
    path = os.path.join(settings.STORAGE_EXPORTS, f"{meeting_id}.docx")

    doc = Document()
    _set_document_default_font(doc, font_name="Nirmala UI", font_size=11)

    title_para = doc.add_heading(meeting.get("title", "Meeting Minutes"), level=1)
    _set_paragraph_font(title_para, font_name="Nirmala UI", font_size=16)

    clean_transcript = (minutes.get("clean_transcript") or "").strip()

    if clean_transcript:
        heading = doc.add_heading(labels["transcript"], level=2)
        _set_paragraph_font(heading, font_name="Nirmala UI", font_size=13)

        para = doc.add_paragraph(clean_transcript)
        _set_paragraph_font(para, font_name="Nirmala UI", font_size=11)

    heading = doc.add_heading(labels["summary"], level=2)
    _set_paragraph_font(heading, font_name="Nirmala UI", font_size=13)

    para = doc.add_paragraph(minutes.get("summary", ""))
    _set_paragraph_font(para, font_name="Nirmala UI", font_size=11)

    heading = doc.add_heading(labels["discussion"], level=2)
    _set_paragraph_font(heading, font_name="Nirmala UI", font_size=13)

    for item in minutes.get("discussion_points", []):
        para = doc.add_paragraph(str(item), style="List Bullet")
        _set_paragraph_font(para, font_name="Nirmala UI", font_size=11)

    heading = doc.add_heading(labels["decisions"], level=2)
    _set_paragraph_font(heading, font_name="Nirmala UI", font_size=13)

    for item in minutes.get("decisions", []):
        para = doc.add_paragraph(str(item), style="List Bullet")
        _set_paragraph_font(para, font_name="Nirmala UI", font_size=11)

    heading = doc.add_heading(labels["actions"], level=2)
    _set_paragraph_font(heading, font_name="Nirmala UI", font_size=13)

    for item in minutes.get("action_items", []):
        owner = item.get("owner_name", "Unassigned")
        task = item.get("task", "")
        deadline = item.get("deadline_text", "")

        text = f"{owner}: {task}"

        if deadline:
            text += f" ({labels['deadline']}: {deadline})"

        para = doc.add_paragraph(text, style="List Bullet")
        _set_paragraph_font(para, font_name="Nirmala UI", font_size=11)

    heading = doc.add_heading(labels["next_steps"], level=2)
    _set_paragraph_font(heading, font_name="Nirmala UI", font_size=13)

    for item in minutes.get("next_steps", []):
        para = doc.add_paragraph(str(item), style="List Bullet")
        _set_paragraph_font(para, font_name="Nirmala UI", font_size=11)

    doc.save(path)

    return path